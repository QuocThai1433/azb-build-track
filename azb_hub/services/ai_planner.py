"""Build a construction plan from a bill of quantities (BOQ) and a schedule workbook, with Claude.

Steps: both workbooks -> compact text (one block per sheet) -> one Claude request with a fixed JSON
schema -> a workbook in the AZB template layout. That workbook then goes through the normal Excel
import (excel_mapping.build_plan -> preview -> confirm), so nothing here touches the database.
"""
from __future__ import annotations

import os
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from ..constants import SUMMARY_TASK_TYPE
from .excel_mapping import is_empty, load_sheets

AI_MODEL = os.environ.get("AZB_AI_MODEL", "claude-opus-5-5")
MAX_INPUT_CHARS = 1_200_000          # about 400k tokens; larger files are refused, never cut silently
MAX_CELL_CHARS = 150


def ai_enabled() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN"))


class AiPlanError(RuntimeError):
    """The plan could not be made. The message is shown to the user as is."""


# ---- the answer Claude must give -------------------------------------------------------------

class AiProject(BaseModel):
    name: str | None = Field(None, description="Tên dự án / công trình")
    client: str | None = Field(None, description="Chủ đầu tư")
    contractor: str | None = Field(None, description="Nhà thầu")
    package_name: str | None = Field(None, description="Gói thầu")
    description: str | None = Field(None, description="Mô tả ngắn các hạng mục chính")


class AiTask(BaseModel):
    id: int = Field(description="Số thứ tự 1, 2, 3… duy nhất, theo đúng thứ tự trong kế hoạch")
    wbs: str | None = Field(None, description="Mã WBS trong file tiến độ, ví dụ 1.2.3")
    category: str | None = Field(None, description="Nhóm hạng mục (cấp cha), ví dụ 'Phần móng'")
    title: str = Field(description="Tên công việc")
    is_summary: bool = Field(description="true nếu là dòng tổng hợp / hạng mục cha")
    planned_days: float | None = Field(None, description="Số ngày thi công")
    start_date: str | None = Field(None, description="YYYY-MM-DD")
    end_date: str | None = Field(None, description="YYYY-MM-DD")
    float_days: float | None = Field(None, description="Dự trữ (ngày), nếu có")
    quantity: float | None = Field(None, description="Khối lượng lấy từ BOQ cho công việc này")
    unit: str | None = Field(None, description="Đơn vị của khối lượng")
    boq_refs: list[str] = Field(default_factory=list, description="STT / mã các dòng BOQ ghép với công việc này")
    notes: str | None = None


class AiCost(BaseModel):
    stt: int
    name: str = Field(description="Nội dung công việc trong BOQ")
    category: str | None = Field(None, description="Nhóm / phần trong BOQ")
    unit: str | None = None
    quantity: float | None = None
    unit_price: float | None = None
    budget_amount: float | None = Field(None, description="Thành tiền")
    notes: str | None = None


class AiResource(BaseModel):
    kind: Literal["Ban chỉ huy", "Công nhật", "Thiết bị"]
    name: str
    quantity: float | None = None
    unit: str | None = None
    duration_months: float | None = None
    start_date: str | None = Field(None, description="YYYY-MM-DD")
    end_date: str | None = Field(None, description="YYYY-MM-DD")
    notes: str | None = None


class AiMethod(BaseModel):
    name: str
    unit: str | None = None
    quantity: float | None = None
    specification: str | None = None
    schedule: str | None = Field(None, description="Thời gian theo tiến độ")


class AiRisk(BaseModel):
    category: str | None = None
    description: str
    basis: str | None = Field(None, description="Căn cứ trong file")
    severity: Literal["Cao", "Trung bình", "Thấp"]
    recommendation: str | None = None


class AiPlan(BaseModel):
    project: AiProject
    tasks: list[AiTask]
    costs: list[AiCost]
    resources: list[AiResource]
    methods: list[AiMethod]
    risks: list[AiRisk]
    warnings: list[str] = Field(description="Những điểm chưa chắc chắn hoặc tự suy ra, người dùng cần kiểm tra")


SYSTEM_PROMPT = """Bạn là kỹ sư lập kế hoạch thi công của nhà thầu xây dựng AZB. Bạn nhận nội dung hai file Excel:
bảng khối lượng (BOQ) và file tiến độ thi công. Hãy dựng bản nháp kế hoạch thi công của dự án theo schema đã cho.

Nguyên tắc:
- File tiến độ là nguồn chính cho danh sách công việc, thứ tự, mã WBS, ngày bắt đầu/kết thúc và số ngày. Giữ đúng
  tên, ngày và cấu trúc cha-con như trong file. Không bỏ công việc nào. Đánh id 1, 2, 3… theo thứ tự xuất hiện.
- Dòng cha / dòng tổng hợp: is_summary = true. category của công việc con là tên hạng mục cha gần nhất.
- Ghép mỗi công việc chi tiết với các dòng BOQ tương ứng: ghi quantity, unit và boq_refs (STT/mã dòng BOQ).
  Nếu một công việc ứng với nhiều dòng BOQ khác đơn vị, chỉ ghi khối lượng chính và nêu phần còn lại trong notes.
- costs: mỗi dòng chi tiết của BOQ (có đơn vị) là một khoản, giữ khối lượng, đơn giá, thành tiền như trong file;
  category là phần/nhóm trong BOQ. Không đưa dòng tổng / dòng tiêu đề vào costs.
- Nếu ngày trong tiến độ bị thiếu, được suy ra từ thứ tự và thời lượng, hoặc BOQ có hạng mục không có trong tiến độ,
  vẫn làm tiếp nhưng ghi rõ vào warnings. Không bịa số liệu: ô không có thông tin thì để null.
- resources, methods, risks: chỉ đề xuất khi có căn cứ trong file (ví dụ khối lượng lớn, tiến độ gấp, thi công song
  song); để mảng rỗng nếu không có căn cứ. Viết tiếng Việt, ngắn gọn.
- Ngày luôn theo dạng YYYY-MM-DD. Số là số thuần (không có dấu phân cách hàng nghìn, không có đơn vị)."""


# ---- workbooks -> text -----------------------------------------------------------------------

def _cell_text(value: Any) -> str:
    if is_empty(value):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = " ".join(str(value).split())
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def workbook_text(path: str | Path, label: str) -> str:
    """Every sheet as tab-separated lines that start with the Excel row number. Empty rows are dropped."""
    parts = [f"===== {label} ====="]
    for name, rows in load_sheets(path).items():
        lines = []
        for number, row in enumerate(rows, 1):
            cells = [_cell_text(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                lines.append(f"{number}\t" + "\t".join(cells))
        if lines:
            parts.append(f"--- Sheet: {name} ---")
            parts.extend(lines)
    return "\n".join(parts)


# ---- the request ----------------------------------------------------------------------------

def request_plan(boq_text: str, schedule_text: str, extra: str = "", client: Any = None) -> tuple[AiPlan, dict[str, Any]]:
    """One Claude request. Returns the plan and a small usage record for the audit log."""
    import anthropic

    if len(boq_text) + len(schedule_text) > MAX_INPUT_CHARS:
        raise AiPlanError("Hai file quá lớn để AI đọc trong một lần. Hãy bỏ bớt các sheet không liên quan "
                          "(bản vẽ, diễn giải chi tiết…) rồi thử lại.")
    note = f"\n\nYêu cầu thêm của người dùng:\n{extra.strip()}" if extra.strip() else ""
    content = (f"{boq_text}\n\n{schedule_text}{note}\n\n"
               "Hãy lập kế hoạch thi công theo schema. Hôm nay là " + date.today().isoformat() + ".")
    client = client or anthropic.Anthropic(max_retries=3, timeout=900)
    try:
        with client.beta.messages.stream(
            model=AI_MODEL,
            max_tokens=64000,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content}],
            thinking={"type": "adaptive"},
            output_config={"effort": "high"},
            output_format=AiPlan,
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
        ) as stream:
            message = stream.get_final_message()
    except anthropic.AuthenticationError as exc:
        raise AiPlanError("API key của Claude không hợp lệ. Hãy kiểm tra biến môi trường ANTHROPIC_API_KEY.") from exc
    except anthropic.RateLimitError as exc:
        raise AiPlanError("Claude đang giới hạn số lượt gọi. Hãy thử lại sau ít phút.") from exc
    except anthropic.BadRequestError as exc:
        raise AiPlanError(f"Claude từ chối yêu cầu: {exc.message}") from exc
    except anthropic.APIStatusError as exc:
        raise AiPlanError(f"Claude báo lỗi {exc.status_code}. Hãy thử lại sau.") from exc
    except anthropic.APIConnectionError as exc:
        raise AiPlanError("Không kết nối được tới Claude. Hãy kiểm tra mạng / proxy của máy chủ.") from exc

    if message.stop_reason == "refusal":
        raise AiPlanError("Claude không xử lý được nội dung hai file này (bị từ chối). Hãy kiểm tra lại file.")
    if message.stop_reason == "max_tokens":
        raise AiPlanError("Kế hoạch quá dài, AI chưa trả lời hết. Hãy tách dự án thành các file nhỏ hơn rồi thử lại.")
    plan = message.parsed_output
    if plan is None:
        raise AiPlanError("AI trả về kết quả không đúng định dạng. Hãy thử lại.")
    usage = {"model": message.model, "input_tokens": message.usage.input_tokens,
             "output_tokens": message.usage.output_tokens}
    return plan, usage


# ---- answer -> workbook in the AZB layout ------------------------------------------------------

TASK_HEADERS = ["ID", "Nhóm hạng mục", "Tên công việc", "Loại", "Số ngày", "Ngày bắt đầu", "Ngày kết thúc",
                "Dự trữ (ngày)", "Ghi chú"]
COST_HEADERS = ["STT", "Nội dung", "Nhóm chi phí", "ĐVT", "Khối lượng", "Đơn giá", "Thành tiền", "Ghi chú"]
RESOURCE_BLOCKS = (("Ban chỉ huy", "Chức danh"), ("Công nhật", "Công việc"), ("Thiết bị", "Hạng mục"))
RESOURCE_HEADERS = ["STT", None, "Số lượng", "Đơn vị", "Thời gian (tháng)", "Bắt đầu", "Kết thúc", "Ghi chú"]
METHOD_HEADERS = ["STT", "Hạng mục", "ĐVT", "Khối lượng", "Quy cách", "Thời gian theo tiến độ", "Nguồn"]
RISK_HEADERS = ["STT", "Nhóm", "Mô tả", "Căn cứ", "Mức độ", "Kiến nghị"]


def _num(value: float | None) -> float | int | None:
    if value is None:
        return None
    return int(value) if float(value).is_integer() else value


def _task_notes(t: AiTask) -> str | None:
    bits = []
    if t.wbs:
        bits.append(f"WBS {t.wbs}")
    if t.quantity is not None:
        bits.append(f"KL: {_num(t.quantity)} {t.unit or ''}".strip())
    if t.boq_refs:
        bits.append("BOQ: " + ", ".join(t.boq_refs))
    if t.notes:
        bits.append(t.notes)
    return " · ".join(bits) or None


def build_workbook(plan: AiPlan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "TongQuan"
    ws.append(["KẾ HOẠCH THI CÔNG (BẢN NHÁP DO AI LẬP)"])
    p = plan.project
    for label, value in (("Tên dự án", p.name), ("Chủ đầu tư", p.client), ("Nhà thầu", p.contractor),
                         ("Gói thầu", p.package_name), ("Mô tả", p.description)):
        ws.append([label, value])

    def sheet(title: str, headers: list[str], rows: list[list[Any]]):
        s = wb.create_sheet(title)
        s.append(headers)
        for row in rows:
            s.append(row)
        return s

    sheet("TienDo", TASK_HEADERS, [
        [t.id, t.category, t.title, SUMMARY_TASK_TYPE if t.is_summary else "Chi tiết", _num(t.planned_days),
         t.start_date, t.end_date, _num(t.float_days), _task_notes(t)] for t in plan.tasks])
    sheet("ChiPhi", COST_HEADERS, [
        [c.stt, c.name, c.category, c.unit, _num(c.quantity), _num(c.unit_price), _num(c.budget_amount), c.notes]
        for c in plan.costs])

    # NguonLuc: one small table per kind; the title of the name column tells the import the kind.
    res = wb.create_sheet("NguonLuc")
    for kind, title_header in RESOURCE_BLOCKS:
        items = [r for r in plan.resources if r.kind == kind]
        if not items:
            continue
        if res.max_row > 1:
            res.append([])
        res.append([h or title_header for h in RESOURCE_HEADERS])
        for i, r in enumerate(items, 1):
            res.append([i, r.name, _num(r.quantity), r.unit, _num(r.duration_months), r.start_date, r.end_date, r.notes])

    sheet("BienPhap", METHOD_HEADERS, [
        [i, m.name, m.unit, _num(m.quantity), m.specification, m.schedule, "AI đề xuất"]
        for i, m in enumerate(plan.methods, 1)])
    sheet("RuiRo", RISK_HEADERS, [
        [i, r.category, r.description, r.basis, r.severity, r.recommendation] for i, r in enumerate(plan.risks, 1)])
    if plan.warnings:
        sheet("GhiChuAI", ["Điểm cần kiểm tra"], [[w] for w in plan.warnings])

    navy = PatternFill("solid", fgColor="172B4D")
    for s in wb.worksheets:
        if s.title == "NguonLuc":
            for row in s.iter_rows():
                if row and row[0].value == "STT":
                    for c in row:
                        c.fill, c.font = navy, Font(color="FFFFFF", bold=True)
        elif s.title != "TongQuan":
            for c in s[1]:
                c.fill, c.font = navy, Font(color="FFFFFF", bold=True)
                c.alignment = Alignment(vertical="center", wrap_text=True)
            s.freeze_panes = "A2"
        for letter, width in zip("ABCDEFGHI", (8, 26, 48, 14, 12, 14, 14, 14, 48)):
            s.column_dimensions[letter].width = width
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 70
    ws["A1"].font = Font(bold=True, size=14)
    for col in (6, 7):
        for row in wb["ChiPhi"].iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = "#,##0"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_plan_workbook(boq_path: str | Path, schedule_path: str | Path, extra: str = "",
                       client: Any = None) -> tuple[bytes, AiPlan, dict[str, Any]]:
    """The whole job: read both files, ask Claude, write the workbook."""
    boq = workbook_text(boq_path, "BẢNG KHỐI LƯỢNG (BOQ)")
    schedule = workbook_text(schedule_path, "FILE TIẾN ĐỘ THI CÔNG")
    plan, usage = request_plan(boq, schedule, extra, client)
    if not plan.tasks:
        raise AiPlanError("AI không tìm thấy công việc nào trong file tiến độ. Hãy kiểm tra lại file.")
    return build_workbook(plan), plan, usage
