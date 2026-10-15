"""Build the multi-sheet report workbook."""
from __future__ import annotations

from io import BytesIO

from sqlalchemy.orm import Session

from .. import repositories as repo

NAVY, WHITE, LIGHT = "172B4D", "FFFFFF", "EAF3F6"

TASK_COLUMNS = ["source_id", "category", "title", "task_type", "planned_days", "start_date", "end_date",
                "plan_progress", "plan_status", "actual_progress", "actual_status", "owner",
                "actual_start_date", "actual_end_date", "float_days", "notes"]
RESOURCE_COLUMNS = ["name", "resource_type", "group_name", "quantity", "duration_months", "start_date", "end_date",
                    "unit", "cost", "notes"]
COST_COLUMNS = ["name", "category", "unit", "quantity", "unit_price", "budget_amount", "actual_amount", "notes"]
RISK_COLUMNS = ["source_id", "category", "description", "basis", "severity", "recommendation", "owner",
                "due_date", "status"]
METHOD_COLUMNS = ["name", "unit", "quantity", "specification", "schedule", "source"]
HSSE_COLUMNS = ["log_date", "log_type", "description", "people_affected", "severity", "status", "reporter"]
DIARY_COLUMNS = ["log_date", "weather", "worker_count", "completed_work", "issues", "reporter"]


def build_report(db: Session, project_id: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project_row = repo.get_project(db, project_id)
    p = repo.to_dict(project_row) if project_row else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "TongQuan"
    ws.append(["BÁO CÁO QUẢN LÝ DỰ ÁN THI CÔNG"])
    ws.merge_cells("A1:D1")
    for label, key in [("Tên dự án", "name"), ("Chủ đầu tư", "client"), ("Nhà thầu", "contractor"),
                       ("Gói thầu", "package_name"), ("Bắt đầu", "start_date"), ("Hoàn thành", "end_date"),
                       ("Ngày báo cáo", "report_date"), ("Giá trị hợp đồng tham chiếu (VND)", "contract_value"),
                       ("Ngân sách Prelim (VND)", "prelim_budget"),
                       ("Tiến độ kế hoạch tại ngày báo cáo", "planned_progress"), ("Ghi chú", "notes")]:
        ws.append([label, p.get(key)])

    def add_sheet(name, headers, rows):
        sheet = wb.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            data = repo.to_dict(row)
            sheet.append([data.get(k) for k in headers])

    add_sheet("TienDo", TASK_COLUMNS, repo.tasks_for_export(db, project_id))
    add_sheet("NguonLuc", RESOURCE_COLUMNS, repo.resources_for_export(db, project_id))
    add_sheet("ChiPhi", COST_COLUMNS, repo.list_costs(db, project_id))
    add_sheet("RuiRo", RISK_COLUMNS, repo.risks_for_export(db, project_id))
    add_sheet("BienPhap", METHOD_COLUMNS, repo.methods_for_export(db, project_id))
    add_sheet("HSSE", HSSE_COLUMNS, repo.safety_logs(db, project_id, limit=None))
    add_sheet("NhatKy", DIARY_COLUMNS, repo.diary_entries(db, project_id))

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.title != "TongQuan":
            for c in sheet[1]:
                c.fill = PatternFill("solid", fgColor=NAVY)
                c.font = Font(color=WHITE, bold=True)
                c.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 32
            sheet.auto_filter.ref = sheet.dimensions
            for row in sheet.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                    if c.row % 2 == 0:
                        c.fill = PatternFill("solid", fgColor="F6F9FC")
            for col in range(1, sheet.max_column + 1):
                lengths = [len(str(sheet.cell(r, col).value or "")) for r in range(1, min(sheet.max_row, 100) + 1)]
                sheet.column_dimensions[get_column_letter(col)].width = min(46, max(12, (max(lengths) if lengths else 10) + 2))
        else:
            for c in sheet[1]:
                c.fill = PatternFill("solid", fgColor=NAVY)
                c.font = Font(color=WHITE, bold=True, size=16)
            for r in range(2, sheet.max_row + 1):
                sheet.cell(r, 1).font = Font(bold=True, color=NAVY)
                sheet.cell(r, 1).fill = PatternFill("solid", fgColor=LIGHT)
            sheet.column_dimensions["A"].width = 38
            sheet.column_dimensions["B"].width = 80
        sheet.sheet_view.showGridLines = False

    for r in range(2, wb["TienDo"].max_row + 1):
        wb["TienDo"].cell(r, 8).number_format = "0.0%"
        wb["TienDo"].cell(r, 10).number_format = "0.0%"
    for r in range(2, wb["ChiPhi"].max_row + 1):
        for col in (5, 6, 7):
            wb["ChiPhi"].cell(r, col).number_format = "#,##0"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


AUDIT_COLUMNS = [("Thời gian", "created_at"), ("Người dùng", "username"), ("Hành động", "action"),
                 ("Đối tượng", "entity_type"), ("Mã bản ghi", "entity_id"), ("Mã dự án", "project_id"),
                 ("Giá trị cũ", "old_value"), ("Giá trị mới", "new_value"), ("IP", "ip_address"),
                 ("Trình duyệt", "user_agent")]


def build_audit_log(rows) -> bytes:
    """One sheet with the audit rows given (already filtered and limited by the caller)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "NhatKyHoatDong"
    ws.append([label for label, _ in AUDIT_COLUMNS])
    for row in rows:
        ws.append([getattr(row, key) for _, key in AUDIT_COLUMNS])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for letter, width in zip("ABCDEFGHIJ", (20, 18, 14, 16, 12, 10, 50, 50, 16, 30)):
        ws.column_dimensions[letter].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
