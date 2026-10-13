"""Read a planning workbook that may not follow the AZB template exactly.

Steps: find which sheet holds which kind of data -> find the header row -> match column titles to
fields -> read the rows. Nothing here touches the database; ``build_plan`` returns a plan that the
preview page shows and that ``excel_import`` writes.

Rules the user chose:
* An empty cell or a missing column stays empty (None). The import still works.
* The user can fix any sheet or column in the preview; aliases the user taught are used first.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

# ---- text helpers ----------------------------------------------------------------------------

_TIMELINE = re.compile(r"^\d{1,2}[/\-.]\d{4}$|^\d{4}[/\-.]\d{1,2}([/\-.]\d{1,2})?$|^\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}$")


def norm(value: Any) -> str:
    """Lower-case, no accents, only letters and digits: 'Hạng mục / công việc' -> 'hangmuccongviec'.
    A header made only of symbols ('#', '%') keeps the symbol."""
    text = "" if value is None else str(value).strip()
    text = text.replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", "", text) or text.strip()


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def col_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _short(value: Any, limit: int = 60) -> str:
    text = value.strftime("%d/%m/%Y") if isinstance(value, (datetime, date)) else str(value).strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


# ---- value conversion (every converter returns None for an empty or unreadable cell) ----------

def to_text(value: Any) -> str | None:
    if is_empty(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def _parse_number(text: str) -> float | None:
    s = text.strip().replace(" ", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if s.count(",") == 1 else s.replace(",", "")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def to_number(value: Any) -> float | None:
    if is_empty(value) or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return _parse_number(str(value))


def to_percent(value: Any) -> float | None:
    """A fraction 0..1. '50%' -> 0.5, 0.5 -> 0.5, 50 -> 0.5."""
    if is_empty(value):
        return None
    if isinstance(value, str) and value.strip().endswith("%"):
        n = _parse_number(value.strip()[:-1])
        return None if n is None else max(0.0, min(1.0, n / 100))
    n = to_number(value)
    if n is None:
        return None
    if n <= 1:
        return max(0.0, n)
    return min(1.0, n / 100)


_DATE_PATTERNS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y")


def to_date(value: Any) -> str | None:
    if is_empty(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= value <= 80000:      # an Excel serial number
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        return None
    text = str(value).strip().split(" ")[0]
    for pattern in _DATE_PATTERNS:
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return None


def to_int(value: Any) -> int | None:
    if is_empty(value) or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else None
    s = str(value).strip()
    return int(s) if s.isdigit() else None


CONVERTERS = {"text": to_text, "number": to_number, "percent": to_percent, "date": to_date, "int": to_int,
              "ordinal": to_text}


# ---- field and entity definitions ------------------------------------------------------------

@dataclass(frozen=True)
class Field:
    key: str
    label: str
    kind: str                       # text | number | percent | date | int | ordinal
    eq: tuple[str, ...] = ()        # normalised header equals one of these (best match)
    has: tuple[Any, ...] = ()       # header contains this text; a tuple means "contains all"
    late: tuple[str, ...] = ()      # equals, but ambiguous: used only when nothing better claims the column


@dataclass(frozen=True)
class Entity:
    key: str
    label: str
    sheet_names: tuple[str, ...]
    fields: tuple[Field, ...]
    title: str
    requires: tuple[str, ...] = ()  # a data row needs these (when most rows have them): skips subtotal rows
    mode: str = "continuous"        # continuous | first_block | blocks
    ignore: tuple[str, ...] = ()    # known headers that carry nothing we store


def _f(key, label, kind, eq=(), has=(), late=()):
    return Field(key, label, kind, tuple(eq), tuple(has), tuple(late))


ENTITIES: dict[str, Entity] = {
    "tasks": Entity("tasks", "Công việc & tiến độ", ("tiendo", "tiendothicong", "ganttwbs", "gantt", "schedule", "progress",
                                                     "congviec", "wbs", "kehoach", "kehoachthicong", "bangtiendo"),
        fields=(
            _f("source_id", "Mã WBS / ID", "int", eq=("id", "wbs", "mawbs", "stt", "sott", "ma", "maso", "no", "#", "mact")),
            _f("category", "Nhóm hạng mục", "text", eq=("nhomhangmuc", "nhom", "nhomcongviec", "khuvuc", "goithau", "goi",
                                                         "category", "group", "hangmucchinh"), has=("nhomhangmuc",), late=("hangmuc",)),
            _f("title", "Tên công việc", "text", eq=("hangmuccongviec", "tencongviec", "congviec", "noidungcongviec", "noidung",
                                                     "tenhangmuc", "taskname", "task", "mota"), late=("hangmuc",)),
            _f("task_type", "Loại", "text", eq=("loai", "loaicv", "loaicongviec", "phanloai", "type", "capdo")),
            _f("planned_days", "Số ngày kế hoạch", "number", eq=("songaymsp", "songay", "thoiluong", "duration", "songaykehoach"),
               has=(("songay", "msp"),)),
            _f("start_date", "Ngày bắt đầu", "date", eq=("batdau", "ngaybatdau", "start", "startdate", "tungay", "ngaybd", "from"),
               has=("ngaybatdau",)),
            _f("end_date", "Ngày kết thúc", "date", eq=("ketthuc", "ngayketthuc", "finish", "end", "enddate", "denngay", "hoanthanh",
                                                         "ngayhoanthanh", "to"), has=("ngayketthuc",)),
            _f("duration_calculated", "Số ngày (tính lại)", "number", eq=("songaytinhlai",), has=(("songay", "tinhlai"),)),
            _f("schedule_check", "Kiểm tra", "text", eq=("kiemtra", "check", "ketquakiemtra")),
            _f("plan_progress", "% kế hoạch", "percent", eq=("tiendokh", "tiendokehoach", "plan", "phantramkehoach", "khkh", "plancomplete"),
               has=("theokh", "tiendokh", "phantramkh", "tiendokehoach")),
            _f("plan_status", "Trạng thái theo KH", "text", eq=("trangthai", "status", "trangthaitheokh", "tinhtrang"),
               has=("trangthai",)),
            _f("tet_flag", "Giao Tết", "text", eq=("giaotet", "tet", "nghitet"), has=("giaotet",)),
            _f("float_days", "Dự trữ (ngày)", "number", eq=("dutru", "float", "slack"), has=("dutru", "float")),
            _f("notes", "Ghi chú", "text", eq=("ghichu", "note", "notes", "remark", "remarks", "nhanxet"), has=("ghichu",)),
        ), title="title", requires=("source_id",)),
    "costs": Entity("costs", "Chi phí", ("chiphi", "chiphiphucvuthicong", "cost", "costs", "budget", "ngansach", "prelim", "dutoan"),
        fields=(
            _f("stt", "STT", "ordinal", eq=("stt", "sott", "id", "ma", "no", "#")),
            _f("name", "Nội dung / khoản chi", "text", eq=("noidungcongviec", "noidung", "tenkhoanchi", "khoanmuc", "tenchiphi",
                                                           "noidungchiphi", "diengiai", "hangmuc", "item", "description", "congviec")),
            _f("unit", "Đơn vị", "text", eq=("dvt", "donvi", "donvitinh", "unit", "dvtinh")),
            _f("quantity", "Khối lượng", "number", eq=("kl", "khoiluong", "soluong", "sl", "qty", "quantity")),
            _f("unit_price", "Đơn giá", "number", eq=("dongia", "dongiavnd", "unitprice", "price"), has=("dongia",)),
            _f("budget_amount", "Thành tiền (dự toán)", "number", eq=("thanhtien", "thanhtienvnd", "dutoan", "ngansach", "sotien",
                                                                      "amount", "budget", "giatri"), has=("thanhtien",)),
            _f("actual_amount", "Thực chi", "number", eq=("thucchi", "thucte", "actual", "dachi", "thuchien")),
            _f("category", "Nhóm chi phí", "text", eq=("nhomchiphi", "loaichiphi", "nhom", "category")),
            _f("notes", "Ghi chú", "text", eq=("ghichu", "note", "notes"), has=("ghichu",)),
        ), title="name", requires=("stt", "unit"), ignore=("tong", "phantramtong", "%tong", "%")),
    "risks": Entity("risks", "Rủi ro & kiến nghị", ("ruiro", "risk", "risks", "ruirokiennghi", "kiennghi", "ruirovakiennghi"),
        fields=(
            _f("source_id", "Mã rủi ro", "int", eq=("#", "stt", "id", "sott", "ma", "no")),
            _f("category", "Nhóm", "text", eq=("nhom", "phanloai", "loai", "category", "linhvuc")),
            _f("description", "Mô tả / phát hiện", "text", eq=("phathien", "mota", "noidung", "ruiro", "vande", "description",
                                                              "noidungruiro", "motaruiro")),
            _f("basis", "Căn cứ", "text", eq=("cancu", "coso", "basis", "bangchung", "nguon")),
            _f("severity", "Mức độ", "text", eq=("mucdo", "mucdoruiro", "severity", "uutien", "priority", "level", "mucdonguyhiem")),
            _f("recommendation", "Kiến nghị", "text", eq=("kiennghi", "dexuat", "recommendation", "giaiphap", "bienphapxuly", "xuly")),
            _f("owner", "Người phụ trách", "text", eq=("nguoiphutrach", "phutrach", "owner", "chutri"), has=("nguoiphutrach",)),
            _f("due_date", "Hạn xử lý", "date", eq=("hanxuly", "han", "duedate", "deadline", "hanhoanthanh")),
        ), title="description", requires=("source_id",)),
    "resources": Entity("resources", "Nguồn lực", ("nguonluc", "resource", "resources", "nhansu", "manpower"),
        fields=(
            _f("stt", "STT", "int", eq=("stt", "sott", "id", "no", "#")),
            _f("name", "Chức danh / hạng mục", "text", eq=("chucdanh", "congviec", "hangmuc", "ten", "tennhansu", "nhansu",
                                                           "thietbi", "noidung", "chucvu", "vitri")),
            _f("quantity", "Số lượng", "number", eq=("soluong", "sl", "qty", "soluongnguoi"), has=("soluong",)),
            _f("duration_months", "Thời gian (tháng)", "number", eq=("sothang", "thoigianthang", "thoigianthangthang"),
               has=(("thoigian", "thang"),)),
            _f("start_date", "Bắt đầu", "date", eq=("batdau", "ngaybatdau", "start", "tungay")),
            _f("end_date", "Kết thúc", "date", eq=("ketthuc", "ngayketthuc", "finish", "denngay")),
            _f("group_name", "Nhóm", "text", eq=("nhom", "bophan", "group")),
            _f("unit", "Đơn vị", "text", eq=("dvt", "donvi", "unit")),
            _f("cost", "Thành tiền", "number", eq=("thanhtien", "thanhtienvnd", "chiphi", "giatri", "cost")),
            _f("note_source", "Nguồn (ghi vào Ghi chú)", "text", eq=("nguon", "source")),
            _f("notes", "Ghi chú", "text", eq=("ghichu", "note", "notes"), has=("ghichu",)),
        ), title="name", requires=("stt",), mode="blocks", ignore=("thoigian",)),
    "methods": Entity("methods", "Biện pháp thi công", ("bienphap", "bienphapthicong", "bienphaptochucthicong", "method", "methods"),
        fields=(
            _f("stt", "STT", "int", eq=("stt", "sott", "id", "no", "#")),
            _f("name", "Hạng mục", "text", eq=("hangmuc", "noidung", "tenhangmuc", "congtrinhtam", "name", "tenbienphap", "bienphap")),
            _f("unit", "Đơn vị", "text", eq=("dvt", "donvi", "unit")),
            _f("quantity", "Khối lượng", "number", eq=("khoiluong", "kl", "soluong", "qty"), has=("khoiluong",)),
            _f("specification", "Quy cách / diễn giải", "text", eq=("quycach", "diengiai", "specification", "mota"), has=("quycach",)),
            _f("schedule", "Thời gian theo tiến độ", "text", eq=("thoigiantheotiendo", "thoigian", "tiendo", "schedule"),
               has=(("thoigian", "tiendo"),)),
            _f("source", "Nguồn", "text", eq=("nguon", "source")),
        ), title="name", requires=("stt",)),
    "hsse": Entity("hsse", "Chỉ tiêu HSSE", ("hsse", "antoan", "antoanlaodong", "an toan", "safety"),
        fields=(
            _f("stt", "STT", "int", eq=("stt", "sott", "id", "no", "#")),
            _f("name", "Chỉ tiêu", "text", eq=("chitieu", "noidung", "hangmuc", "tenchitieu", "indicator")),
            _f("value", "Giá trị", "number", eq=("giatri", "sl", "soluong", "value")),
            _f("unit", "Đơn vị", "text", eq=("donvi", "dvt", "unit")),
            _f("notes", "Ghi chú", "text", eq=("ghichu", "note", "notes"), has=("ghichu",)),
        ), title="name", requires=("stt",), mode="first_block"),
}
ENTITY_ORDER = ("tasks", "costs", "risks", "resources", "methods", "hsse")
OVERVIEW_SHEETS = ("tongquan", "overview", "thongtinchung", "thongtin", "tongquat", "tonghop")

PROJECT_FIELDS = (
    # key, label, kind, labels that equal the cell, labels contained in the cell, fallback (row, column) 1-based in the AZB template
    ("name", "Tên dự án", "text", ("duan", "tenduan", "tencongtrinh", "congtrinh", "project", "projectname"), (), (4, 3)),
    ("client", "Chủ đầu tư", "text", ("chudautu", "cdt", "client", "owner"), ("khachhang", "chudautu"), (5, 3)),
    ("contractor", "Nhà thầu", "text", ("nhathau", "contractor"), ("nhathau",), (6, 3)),
    ("package_name", "Gói thầu", "text", ("goithau", "package"), ("goithau",), (7, 3)),
    ("description", "Mô tả / hạng mục chính", "text", ("mota", "hangmucchinh", "description"), ("hangmucchinh",), (8, 3)),
    ("report_date", "Ngày báo cáo", "date", ("ngaybaocao",), ("ngaybaocao",), None),
    ("prelim_budget", "Ngân sách Prelim", "number", (), (("tongchiphi", "prelim"),), (25, 3)),
    ("planned_progress", "% tiến độ kế hoạch tại ngày báo cáo", "percent", (), (("thoigian", "troiqua"), ("tiendo", "kehoach")), (16, 3)),
)


# ---- plan objects ----------------------------------------------------------------------------

@dataclass
class ColumnPlan:
    index: int
    letter: str
    header: str
    norm: str
    field: str | None
    origin: str            # auto | alias | manual | none
    ignorable: bool
    samples: list[str] = field(default_factory=list)


@dataclass
class EntityPlan:
    entity: str
    label: str
    sheet: str | None = None
    sheet_origin: str = "none"     # name | alias | content | manual | none
    layout: str = "table"          # table | azb_resources
    header_row: int | None = None  # 1-based
    columns: list[ColumnPlan] = field(default_factory=list)
    rows: list[dict[str, Any]] = field(default_factory=list)
    skipped: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    mapped_fields: set[str] = field(default_factory=set)

    @property
    def used(self) -> bool:
        return bool(self.rows)


@dataclass
class ImportPlan:
    file_name: str
    sheet_names: list[str]
    project: dict[str, Any] = field(default_factory=dict)        # only the facts found in the file
    project_sources: dict[str, str] = field(default_factory=dict)
    entities: dict[str, EntityPlan] = field(default_factory=dict)
    unused_sheets: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def has_data(self) -> bool:
        return any(e.used for e in self.entities.values())


class WorkbookReadError(ValueError):
    """The file could not be opened as an .xlsx workbook."""


def load_sheets(path: str | Path) -> dict[str, list[tuple]]:
    """Every sheet as a list of row tuples (cached values, no formulas)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:   # pragma: no cover
        raise RuntimeError("Chưa cài openpyxl. Chạy: pip install -r requirements.txt") from exc
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookReadError(f"Không mở được file như một file Excel (.xlsx). Chi tiết: {exc}") from exc
    try:
        return {name: [tuple(r) for r in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}
    finally:
        wb.close()


# ---- matching --------------------------------------------------------------------------------

def _match(f: Field, n: str) -> tuple[int, int] | None:
    if n in f.eq:
        return (0, f.eq.index(n))
    for i, item in enumerate(f.has):
        parts = (item,) if isinstance(item, str) else item
        if all(p in n for p in parts):
            return (1, i)
    if n in f.late:
        return (2, f.late.index(n))
    return None


def _assign(entity: Entity, header: tuple, aliases: dict, manual: dict[int, str] | None):
    """Match the cells of one header row to fields. Returns {column index: (field key | None, origin)}."""
    assigned: dict[int, tuple[str | None, str]] = {}
    candidates: list[tuple[tuple[int, int], int, str]] = []
    for col, raw in enumerate(header):
        if is_empty(raw) or isinstance(raw, (datetime, date)) or _TIMELINE.match(str(raw).strip()):
            continue
        n = norm(raw)
        if n in aliases.get(("column", entity.key), {}):
            target = aliases[("column", entity.key)][n]
            assigned[col] = (target or None, "alias")
            continue
        for f in entity.fields:
            p = _match(f, n)
            if p is not None:
                candidates.append((p, col, f.key))
    taken_cols = set(assigned)
    taken_fields = {v[0] for v in assigned.values() if v[0]}
    for _p, col, key in sorted(candidates):
        if col in taken_cols or key in taken_fields:
            continue
        assigned[col] = (key, "auto")
        taken_cols.add(col)
        taken_fields.add(key)
    if manual:
        for col, key in manual.items():
            if key:
                for other, (k, _o) in list(assigned.items()):
                    if k == key and other != col:
                        assigned[other] = (None, "none")
            assigned[col] = (key or None, "manual")
    return assigned


def _score(entity: Entity, assigned: dict[int, tuple[str | None, str]]) -> int:
    keys = {k for k, _o in assigned.values() if k}
    return len(keys) if entity.title in keys else 0


def _find_header_rows(entity: Entity, rows: list[tuple], aliases: dict, limit: int | None,
                      min_score: int = 3) -> list[int]:
    """0-based indexes of rows that look like this entity's header row.

    min_score is how many known columns a row needs (the title column is always one of them).
    It is 3 when we only have the content to go on, and lower when the sheet is already known by name,
    by a remembered choice or by the user, so a file with few columns still imports.
    """
    scored: list[tuple[int, int]] = []
    for i, row in enumerate(rows if limit is None else rows[:limit]):
        if not any(isinstance(v, str) and v.strip() for v in row):
            continue
        score = _score(entity, _assign(entity, row, aliases, None))
        if score >= min_score:
            scored.append((i, score))
    if not scored:
        return []
    if entity.mode == "blocks":
        return [i for i, _s in scored]
    best = max(s for _i, s in scored)
    return [next(i for i, s in scored if s == best)]


# ---- workbook layout of the AZB template for NguonLuc (three small tables in one sheet) --------

def _azb_resources(rows: list[tuple]) -> list[dict[str, Any]]:
    def cell(r, c):
        return rows[r - 1][c - 1] if r - 1 < len(rows) and c - 1 < len(rows[r - 1]) else None

    out: list[dict[str, Any]] = []
    for rn in range(1, len(rows) + 1):
        name = to_text(cell(rn, 2))
        if 5 <= rn <= 17 and name and cell(rn, 3) is not None:
            out.append({"name": name, "resource_type": "Ban chỉ huy", "group_name": to_text(cell(rn, 7)),
                        "quantity": to_number(cell(rn, 3)), "duration_months": to_number(cell(rn, 4)),
                        "start_date": to_date(cell(rn, 5)), "end_date": to_date(cell(rn, 6)), "unit": "người",
                        "cost": None, "notes": to_text(cell(rn, 8))})
        elif 33 <= rn <= 35 and name and cell(rn, 3) is not None:
            out.append({"name": name, "resource_type": "Công nhật", "group_name": "Công nhật",
                        "quantity": to_number(cell(rn, 3)), "duration_months": to_number(cell(rn, 4)),
                        "start_date": to_date(cell(rn, 5)), "end_date": to_date(cell(rn, 6)), "unit": "người",
                        "cost": None, "notes": to_text(cell(rn, 8))})
        elif 40 <= rn <= 46 and name and cell(rn, 4) is not None:
            notes = " ".join(x for x in (to_text(cell(rn, 8)), to_text(cell(rn, 7))) if x)
            out.append({"name": name, "resource_type": "Thiết bị / vận chuyển", "group_name": "Thiết bị",
                        "quantity": to_number(cell(rn, 4)), "duration_months": None, "start_date": None,
                        "end_date": None, "unit": to_text(cell(rn, 3)), "cost": to_number(cell(rn, 5)),
                        "notes": notes or None})
    return out


_RESOURCE_TYPE_BY_TITLE = {"chucdanh": "Ban chỉ huy", "congviec": "Công nhật", "hangmuc": "Thiết bị / vận chuyển"}


# ---- reading rows ----------------------------------------------------------------------------

def _read_block(entity: Entity, plan: EntityPlan, rows: list[tuple], header_idx: int, end_idx: int,
                colmap: dict[int, str], block_type: str | None = None) -> list[dict[str, Any]]:
    kinds = {f.key: f.kind for f in entity.fields}
    candidates: list[tuple[int, dict[str, Any]]] = []
    bad: dict[str, int] = {}
    for r in range(header_idx + 1, end_idx):
        row = rows[r]
        values: dict[str, Any] = {}
        for col, key in colmap.items():
            raw = row[col] if col < len(row) else None
            if is_empty(raw):
                values[key] = None
                continue
            converted = CONVERTERS[kinds[key]](raw)
            if converted is None and kinds[key] not in ("text", "ordinal") and key != "stt":
                bad[key] = bad.get(key, 0) + 1
            values[key] = converted
        raw_any = [row[c] for c in colmap if c < len(row) and not is_empty(row[c])]
        if not raw_any:
            continue
        if values.get(entity.title) is None:
            plan.skipped.append((r + 1, "không có " + next(f.label for f in entity.fields if f.key == entity.title).lower()))
            continue
        if block_type:
            values["resource_type"] = block_type
        values["_row"] = r + 1
        candidates.append((r + 1, values))

    # A "detail row" rule (skip subtotal / heading rows) is used only when most rows satisfy it.
    required = [k for k in entity.requires if k in colmap.values()]
    active = []
    for key in required:
        covered = sum(1 for _r, v in candidates if v.get(key) is not None)
        if candidates and covered / len(candidates) >= 0.5:
            active.append(key)
    kept: list[dict[str, Any]] = []
    for rownum, values in candidates:
        missing = next((k for k in active if values.get(k) is None), None)
        if missing:
            label = next(f.label for f in entity.fields if f.key == missing)
            plan.skipped.append((rownum, f"không có {label.lower()} (có vẻ là dòng tiêu đề / dòng tổng)"))
            continue
        kept.append(values)
    for key, count in bad.items():
        label = next(f.label for f in entity.fields if f.key == key)
        plan.warnings.append(f"{count} ô “{label}” không đọc được nên để trống.")
    return kept


def _build_entity(entity: Entity, sheet_name: str, origin: str, rows: list[tuple], aliases: dict,
                  manual: dict[int, str] | None) -> EntityPlan:
    plan = EntityPlan(entity.key, entity.label, sheet=sheet_name, sheet_origin=origin)

    if entity.key == "resources" and not manual and norm(sheet_name) == "nguonluc" and len(rows) >= 5 \
            and norm(rows[3][1] if len(rows[3]) > 1 else None) == "chucdanh":
        plan.layout = "azb_resources"
        plan.header_row = 4
        plan.rows = _azb_resources(rows)
        plan.mapped_fields = {f.key for f in entity.fields}
        return plan

    known = origin in ("name", "alias", "manual")
    minimum = (2 if known else 3) if entity.mode == "blocks" else (1 if known else 3)
    headers = _find_header_rows(entity, rows, aliases, None if entity.mode == "blocks" else 40, minimum)
    if not headers and manual:
        headers = [0]
    if not headers:
        return plan
    plan.header_row = headers[0] + 1
    first = _assign(entity, rows[headers[0]], aliases, manual)
    plan.mapped_fields = {k for k, _o in first.values() if k}

    # Columns shown in the preview: every titled column of the first header row.
    header_cells = rows[headers[0]]
    data_rows = rows[headers[0] + 1:]
    for col, raw in enumerate(header_cells):
        if is_empty(raw) or isinstance(raw, (datetime, date)) or _TIMELINE.match(str(raw).strip()):
            continue
        key, how = first.get(col, (None, "none"))
        samples = []
        for r in data_rows:
            if col < len(r) and not is_empty(r[col]):
                samples.append(_short(r[col], 28))
                if len(samples) == 2:
                    break
        n = norm(raw)
        plan.columns.append(ColumnPlan(col, col_letter(col), str(raw).strip(), n, key, how if key or how != "auto" else "none",
                                       n in entity.ignore, samples))

    if entity.mode == "blocks":
        for k, h in enumerate(headers):
            end = headers[k + 1] if k + 1 < len(headers) else len(rows)
            for r in range(h + 1, end):                     # a block ends at its first empty row
                if all(is_empty(v) for v in rows[r]):
                    end = r
                    break
            assigned = _assign(entity, rows[h], aliases, manual if k == 0 else None)
            colmap = {c: key for c, (key, _o) in assigned.items() if key}
            title_col = next((c for c, key in colmap.items() if key == entity.title), None)
            btype = None
            if entity.key == "resources":
                btype = _RESOURCE_TYPE_BY_TITLE.get(norm(rows[h][title_col]) if title_col is not None else "", "Nhân sự")
            plan.rows += _read_block(entity, plan, rows, h, end, colmap, btype)
    else:
        end = len(rows)
        if entity.mode == "first_block":
            seen = False
            for r in range(headers[0] + 1, len(rows)):
                empty = all(is_empty(v) for v in rows[r])
                if seen and empty:
                    end = r
                    break
                seen = seen or not empty
        colmap = {c: key for c, (key, _o) in first.items() if key}
        plan.rows = _read_block(entity, plan, rows, headers[0], end, colmap)
    return plan


# ---- project facts ---------------------------------------------------------------------------

def _right_of(row: tuple, col: int) -> tuple[int, Any] | None:
    for c in range(col + 1, min(len(row), col + 5)):
        if not is_empty(row[c]):
            return c, row[c]
    return None


def _extract_project(sheets: dict[str, list[tuple]], overview: str | None) -> tuple[dict, dict]:
    values: dict[str, Any] = {}
    sources: dict[str, str] = {}
    search = [s for s in (overview, next((n for n in sheets if norm(n) == "tiendo"), None)) if s]
    for key, _label, kind, eq, has, _pos in PROJECT_FIELDS:
        found = False
        for sheet in search:
            for r, row in enumerate(sheets[sheet][:60], 1):
                for c, cell in enumerate(row[:4]):
                    if not isinstance(cell, str) or not cell.strip():
                        continue
                    n = norm(cell)
                    ok = n in eq or any(all(p in n for p in ((h,) if isinstance(h, str) else h)) for h in has)
                    if not ok or len(n) > 90:
                        continue
                    hit = _right_of(row, c)
                    if hit is None:
                        continue
                    converted = CONVERTERS[kind](hit[1])
                    if converted is None:
                        continue
                    values[key] = converted
                    sources[key] = f"{sheet}!{col_letter(hit[0])}{r}"
                    found = True
                    break
                if found:
                    break
            if found:
                break
        if not found and overview and norm(overview) == "tongquan":
            pos = next(p for k, *_rest, p in PROJECT_FIELDS if k == key)
            if pos:
                rows = sheets[overview]
                r, c = pos
                if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
                    converted = CONVERTERS[kind](rows[r - 1][c - 1])
                    if converted is not None:
                        values[key] = converted
                        sources[key] = f"{overview}!{col_letter(c - 1)}{r} (vị trí mẫu AZB)"
    return values, sources


# ---- the plan --------------------------------------------------------------------------------

def build_plan(sheets: dict[str, list[tuple]], file_name: str, aliases: dict | None = None,
               overrides: dict | None = None) -> ImportPlan:
    """aliases: {("column", entity): {norm header: field or ""}, ("sheet", entity): {norm sheet: ""}}.
    overrides: {entity: {"sheet": name | "" | None, "columns": {column index: field | ""}}} from the preview form."""
    aliases = aliases or {}
    overrides = overrides or {}
    plan = ImportPlan(file_name=file_name, sheet_names=list(sheets))
    assigned_sheets: dict[str, tuple[str, str]] = {}      # entity -> (sheet, origin)

    overview = next((n for n in sheets if norm(n) in OVERVIEW_SHEETS), None)
    taken = {overview} if overview else set()

    # 1. sheets the user chose
    for key in ENTITY_ORDER:
        chosen = (overrides.get(key) or {}).get("sheet")
        if chosen is not None:
            if chosen and chosen in sheets:
                assigned_sheets[key] = (chosen, "manual")
                taken.add(chosen)
            else:
                assigned_sheets[key] = ("", "manual")      # "not used"
    # 2. sheets the user taught before, then the sheet's own name
    for key in ENTITY_ORDER:
        if key in assigned_sheets:
            continue
        learned = aliases.get(("sheet", key), {})
        pick = next((n for n in sheets if n not in taken and norm(n) in learned), None)
        if pick:
            assigned_sheets[key] = (pick, "alias")
            taken.add(pick)
    for exact in (True, False):
        for key in ENTITY_ORDER:
            if key in assigned_sheets:
                continue
            names = ENTITIES[key].sheet_names
            pick = next((n for n in sheets if n not in taken and (norm(n) in names if exact else
                         any(len(s) >= 5 and s in norm(n) for s in names))), None)
            if pick:
                assigned_sheets[key] = (pick, "name")
                taken.add(pick)
    # 3. by content: a sheet whose header row looks like the entity
    for key in ENTITY_ORDER:
        if key in assigned_sheets:
            continue
        best = None
        for name, rows in sheets.items():
            if name in taken:
                continue
            for i in _find_header_rows(ENTITIES[key], rows[:40], aliases, 40):
                score = _score(ENTITIES[key], _assign(ENTITIES[key], rows[i], aliases, None))
                if best is None or score > best[0]:
                    best = (score, name)
        if best:
            assigned_sheets[key] = (best[1], "content")
            taken.add(best[1])

    # 4. read every entity
    for key in ENTITY_ORDER:
        entity = ENTITIES[key]
        sheet, origin = assigned_sheets.get(key, ("", "none"))
        if not sheet:
            plan.entities[key] = EntityPlan(key, entity.label, sheet_origin="none")
            continue
        manual = (overrides.get(key) or {}).get("columns") or None
        plan.entities[key] = _build_entity(entity, sheet, origin, sheets[sheet], aliases, manual)

    plan.project, plan.project_sources = _extract_project(sheets, overview)
    used_sheets = {e.sheet for e in plan.entities.values() if e.sheet and e.used}
    if overview:
        used_sheets.add(overview)
    plan.unused_sheets = [n for n in sheets if n not in used_sheets]
    return plan
