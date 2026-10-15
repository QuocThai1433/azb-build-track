"""Sign in, roles, project membership, cost approval and the audit log."""
import re
import sqlite3

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from azb_hub import auth
from azb_hub import repositories as repo
from azb_hub.db import DB_PATH, SessionLocal
from azb_hub.models import AuditLog, Cost, Project, Task, User

SLUG1 = "nha-may-san-xuat-cac-san-pham-sau-thep"


def _make_user(username: str, role: str, password: str = "Mat-khau-123", projects=(), must_change=False) -> User:
    with SessionLocal() as db:
        user = repo.get_user_by_username(db, username)
        if user is None:
            user = User(username=username, full_name=username.title(), role=role, must_change_password=must_change,
                        password_hash=auth.hash_password(password), is_active=True)
            db.add(user)
            db.flush()
        repo.set_memberships(db, user.id, set(projects))
        db.commit()
        return user


def _login(username: str, password: str = "Mat-khau-123") -> TestClient:
    from azb_hub.main import app
    c = TestClient(app, follow_redirects=False)
    r = c.post("/login", data={"username": username, "password": password})
    assert r.status_code == 303, r.text[:200]
    return c


@pytest.fixture()
def second_project(client):
    """A project the staff member is NOT assigned to."""
    r = client.post("/projects/new", data={"name": "Dự án khác của giám đốc"})
    slug = re.match(r"^/([a-z0-9-]+)/$", r.headers["location"]).group(1)
    with SessionLocal() as db:
        return db.scalar(select(Project).where(Project.slug == slug))


@pytest.fixture()
def staff():
    _make_user("nv-test", auth.NHAN_VIEN, projects={1})
    return _login("nv-test")


def _audit_count(action: str, **where) -> int:
    with SessionLocal() as db:
        stmt = select(func.count()).select_from(AuditLog).where(AuditLog.action == action)
        for key, value in where.items():
            stmt = stmt.where(getattr(AuditLog, key) == value)
        return db.scalar(stmt)


# ---- sign in ----------------------------------------------------------------------------------

def test_anonymous_is_sent_to_login(client):
    from azb_hub.main import app
    anon = TestClient(app, follow_redirects=False)
    r = anon.get(f"/{SLUG1}/tasks?q=nền")
    assert r.status_code == 303
    assert r.headers["location"].startswith("/login?next=")
    assert anon.get("/login").status_code == 200
    assert anon.get("/api/projects/search").status_code == 401
    assert anon.post(f"/{SLUG1}/tasks/1/delete").status_code == 303        # no action without signing in


def test_wrong_password_is_refused_and_logged(client):
    before = _audit_count("login_failed")
    from azb_hub.main import app
    c = TestClient(app, follow_redirects=False)
    r = c.post("/login", data={"username": "admin", "password": "sai-mat-khau"})
    assert r.status_code == 200 and "Sai tên đăng nhập hoặc mật khẩu" in r.text
    assert _audit_count("login_failed") == before + 1
    assert c.get(f"/{SLUG1}/").status_code == 303                            # still anonymous


def test_account_locks_after_repeated_failures(client):
    _make_user("hay-quen", auth.NHAN_VIEN)
    from azb_hub.main import app
    c = TestClient(app, follow_redirects=False)
    for _ in range(5):
        c.post("/login", data={"username": "hay-quen", "password": "sai"})
    r = c.post("/login", data={"username": "hay-quen", "password": "Mat-khau-123"})     # right password, but locked
    assert r.status_code == 200 and "tạm khóa" in r.text


def test_open_redirect_is_blocked(client):
    _make_user("nv-redirect", auth.NHAN_VIEN)
    from azb_hub.main import app
    c = TestClient(app, follow_redirects=False)
    r = c.post("/login", data={"username": "nv-redirect", "password": "Mat-khau-123", "next": "//evil.test/x"})
    assert r.headers["location"] == "/"


def test_inactive_user_cannot_sign_in(client):
    _make_user("da-khoa", auth.NHAN_VIEN)
    with SessionLocal() as db:
        repo.get_user_by_username(db, "da-khoa").is_active = False
        db.commit()
    from azb_hub.main import app
    r = TestClient(app, follow_redirects=False).post("/login", data={"username": "da-khoa", "password": "Mat-khau-123"})
    assert r.status_code == 200


def test_first_login_forces_password_change(client):
    _make_user("moi-vao", auth.NHAN_VIEN, must_change=True, projects={1})
    c = _login("moi-vao")
    r = c.get(f"/{SLUG1}/tasks")
    assert r.status_code == 303 and r.headers["location"] == "/account/password"
    r = c.post("/account/password", data={"current": "Mat-khau-123", "new": "Mat-khau-moi-456", "confirm": "Mat-khau-moi-456"})
    assert r.status_code == 303
    assert c.get(f"/{SLUG1}/tasks").status_code == 200


def test_logout_ends_the_session(client):
    _make_user("nv-out", auth.NHAN_VIEN, projects={1})
    c = _login("nv-out")
    assert c.get(f"/{SLUG1}/tasks").status_code == 200
    assert c.post("/logout").status_code == 303
    assert c.get(f"/{SLUG1}/tasks").status_code == 303


def test_password_is_not_stored_in_clear():
    with SessionLocal() as db:
        stored = repo.get_user_by_username(db, "admin").password_hash
    assert stored.startswith("scrypt$") and "Admin-test-123" not in stored
    assert auth.verify_password("Admin-test-123", stored) and not auth.verify_password("khac", stored)


# ---- roles and membership ---------------------------------------------------------------------

def test_staff_gets_403_on_unassigned_project(client, staff, second_project):
    assert staff.get(f"/{SLUG1}/tasks").status_code == 200
    assert staff.get(f"/{second_project.slug}/").status_code == 403
    assert staff.get(f"/{second_project.slug}/tasks").status_code == 403
    assert _audit_count("denied") >= 1


def test_staff_sees_only_assigned_projects(client, staff, second_project):
    page = staff.get("/projects").text
    assert "Dự án khác của giám đốc" not in page
    assert "Dự án khác của giám đốc" in client.get("/projects").text


def test_staff_cannot_delete_records(client, staff):
    r = client.post(f"/{SLUG1}/tasks/new", data={"title": "Việc để thử xóa"})
    with SessionLocal() as db:
        task_id = db.scalar(select(Task.id).where(Task.title == "Việc để thử xóa"))
    assert staff.post(f"/{SLUG1}/tasks/{task_id}/delete").status_code == 403
    for path in ("costs/1", "risks/1", "resources/1", "methods/1", "hsse/1", "diary/1"):
        assert staff.post(f"/{SLUG1}/{path}/delete").status_code == 403
    with SessionLocal() as db:
        assert db.get(Task, task_id) is not None
    assert client.post(f"/{SLUG1}/tasks/{task_id}/delete").status_code == 303            # the director may


def test_staff_cannot_manage_projects_import_export_or_users(client, staff, second_project):
    assert staff.get("/projects/new").status_code == 403
    assert staff.post("/projects/new", data={"name": "Lén tạo"}).status_code == 403
    assert staff.post(f"/{SLUG1}/archive").status_code == 403
    assert staff.post(f"/{SLUG1}/delete").status_code == 403
    assert staff.get(f"/{SLUG1}/import").status_code == 403
    assert staff.get(f"/{SLUG1}/reports/export.xlsx").status_code == 403
    assert staff.post(f"/{SLUG1}/project", data={"name": "Đổi tên"}).status_code == 403
    assert staff.get("/users").status_code == 403
    assert client.get("/users").status_code == 200


def test_staff_can_enter_data_in_assigned_project(client, staff):
    r = staff.post(f"/{SLUG1}/risks/new", data={"description": "Mưa lớn trễ đổ bê tông", "severity": "Cao"})
    assert r.status_code == 303
    page = staff.get(f"/{SLUG1}/risks").text
    assert "Mưa lớn trễ đổ bê tông" in page
    assert not re.search(r"/risks/\d+/delete", page)                      # listed, but no delete button


def test_hidden_buttons_follow_the_role(client, staff):
    assert re.search(r"/costs/\d+/delete", client.get(f"/{SLUG1}/costs").text)
    assert not re.search(r"/costs/\d+/delete", staff.get(f"/{SLUG1}/costs").text)
    assert "Cập nhật từ Excel" not in staff.get(f"/{SLUG1}/").text


# ---- cost approval ----------------------------------------------------------------------------

def test_staff_cost_waits_for_approval_then_director_approves(client, staff):
    staff.post(f"/{SLUG1}/costs/new", data={"name": "Thuê máy xúc chờ duyệt", "actual_amount": "5000000"})
    with SessionLocal() as db:
        cost = db.scalar(select(Cost).where(Cost.name == "Thuê máy xúc chờ duyệt"))
        cost_id = cost.id
        assert cost.approval_status == "cho_duyet"
    assert staff.post(f"/{SLUG1}/costs/{cost_id}/approval", data={"decision": "approve"}).status_code == 403
    before = _audit_count("approve")
    assert client.post(f"/{SLUG1}/costs/{cost_id}/approval", data={"decision": "approve"}).status_code == 303
    with SessionLocal() as db:
        cost = db.get(Cost, cost_id)
        assert cost.approval_status == "da_duyet" and cost.approved_by and cost.approved_at
    assert _audit_count("approve") == before + 1
    staff.post(f"/{SLUG1}/costs/{cost_id}/edit", data={"name": "Thuê máy xúc chờ duyệt", "actual_amount": "9000000"})
    with SessionLocal() as db:
        assert db.get(Cost, cost_id).approval_status == "cho_duyet"                   # editing asks for approval again


def test_director_cost_is_approved_at_once(client):
    client.post(f"/{SLUG1}/costs/new", data={"name": "Chi phí của giám đốc"})
    with SessionLocal() as db:
        assert db.scalar(select(Cost.approval_status).where(Cost.name == "Chi phí của giám đốc")) == "da_duyet"


# ---- audit log --------------------------------------------------------------------------------

def test_create_update_delete_are_logged_with_the_user(client):
    client.post(f"/{SLUG1}/tasks/new", data={"title": "Việc có nhật ký", "category": "Móng"})
    with SessionLocal() as db:
        task_id = db.scalar(select(Task.id).where(Task.title == "Việc có nhật ký"))
        created = db.scalar(select(AuditLog).where(AuditLog.action == "create", AuditLog.entity_type == "tasks",
                                                   AuditLog.entity_id == task_id))
    assert created is not None and created.username == "admin" and created.project_id == 1
    assert "Việc có nhật ký" in created.new_value and created.ip_address

    client.post(f"/{SLUG1}/tasks/{task_id}/edit", data={"title": "Việc có nhật ký", "category": "Móng", "owner": "An"})
    with SessionLocal() as db:
        updated = db.scalar(select(AuditLog).where(AuditLog.action == "update", AuditLog.entity_id == task_id,
                                                   AuditLog.entity_type == "tasks"))
    assert updated is not None and '"owner": "An"' in updated.new_value and "owner" in updated.old_value

    client.post(f"/{SLUG1}/tasks/{task_id}/delete")
    with SessionLocal() as db:
        deleted = db.scalar(select(AuditLog).where(AuditLog.action == "delete", AuditLog.entity_id == task_id,
                                                   AuditLog.entity_type == "tasks"))
    assert deleted is not None and "Việc có nhật ký" in deleted.old_value


def test_login_logout_and_export_are_logged(client):
    before = _audit_count("export")
    assert client.get(f"/{SLUG1}/reports/export.xlsx").status_code == 200
    assert _audit_count("export") == before + 1
    _make_user("nv-log", auth.NHAN_VIEN)
    c = _login("nv-log")
    c.post("/logout")
    assert _audit_count("login", username="nv-log") >= 1 and _audit_count("logout", username="nv-log") >= 1


def test_audit_log_cannot_be_changed_or_deleted(client):
    con = sqlite3.connect(DB_PATH)
    try:
        with pytest.raises(sqlite3.DatabaseError):
            con.execute("UPDATE audit_logs SET action = 'x'")
        with pytest.raises(sqlite3.DatabaseError):
            con.execute("DELETE FROM audit_logs")
    finally:
        con.close()


def test_project_delete_keeps_its_history(client, second_project):
    pid = second_project.id
    assert _audit_count("create", entity_type="project", project_id=pid) == 1
    client.post(f"/{second_project.slug}/delete")
    assert _audit_count("delete", entity_type="project", project_id=pid) == 1
    assert _audit_count("create", entity_type="project", project_id=pid) == 1          # history survives


# ---- audit log page ---------------------------------------------------------------------------

def test_audit_page_is_for_director_and_admin_only(client, staff):
    assert client.get("/audit-logs").status_code == 200
    assert staff.get("/audit-logs").status_code == 403
    assert staff.get("/audit-logs/export.xlsx").status_code == 403
    assert "Nhật ký hoạt động" not in staff.get(f"/{SLUG1}/").text                 # no menu entry for staff
    _make_user("gd-test", auth.GIAM_DOC)
    assert _login("gd-test").get("/audit-logs").status_code == 200


def test_audit_page_filters_and_pages(client):
    client.post(f"/{SLUG1}/tasks/new", data={"title": "Việc để lọc nhật ký"})
    page = client.get("/audit-logs?action=create&project_id=1").text
    assert "Việc để lọc nhật ký" in page
    assert "Việc để lọc nhật ký" not in client.get("/audit-logs?action=login").text
    assert client.get("/audit-logs?page=9999").status_code == 200                    # clamped to the last page
    assert client.get("/audit-logs?user_id=abc&date_from=zzz").status_code == 200   # junk input never breaks it
    future = client.get("/audit-logs?date_from=2999-01-01").text
    assert "Không có bản ghi nào khớp bộ lọc" in future


def test_audit_export_is_a_workbook_and_is_logged(client):
    from io import BytesIO

    from openpyxl import load_workbook
    before = _audit_count("export", entity_type="audit_logs")
    r = client.get("/audit-logs/export.xlsx?action=create")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:3] == ["Thời gian", "Người dùng", "Hành động"]
    assert ws.max_row > 1 and all(row[2].value == "create" for row in ws.iter_rows(min_row=2))
    assert _audit_count("export", entity_type="audit_logs") == before + 1


def test_admin_can_delete_user_but_not_self_or_last_admin(client):
    victim = _make_user("xoa_toi", "nhan_vien")
    r = client.post(f"/users/{victim.id}/delete")
    assert r.status_code == 303
    with SessionLocal() as db:
        assert repo.get_user(db, victim.id) is None
        assert any(l.action == "delete" and l.entity_type == "users" and l.entity_id == victim.id
                   for l in repo.list_audit_logs(db, limit=20))
        admin = repo.get_user_by_username(db, "admin")
    client.post(f"/users/{admin.id}/delete")                    # cannot delete yourself
    with SessionLocal() as db:
        assert repo.get_user(db, admin.id) is not None
    assert client.post("/users/999999/delete").status_code == 404


def test_staff_cannot_delete_users(client):
    victim = _make_user("khong_xoa", "nhan_vien")
    _make_user("nv_khong_xoa", "nhan_vien")
    c = _login("nv_khong_xoa")
    assert c.post(f"/users/{victim.id}/delete").status_code == 403
    with SessionLocal() as db:
        assert repo.get_user(db, victim.id) is not None
