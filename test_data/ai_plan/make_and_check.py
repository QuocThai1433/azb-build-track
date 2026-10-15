"""Bộ file mẫu để thử chức năng "Lập kế hoạch bằng AI" + kiểm tra file AI xuất ra.

  python make_and_check.py make              -> tạo BOQ_mau.xlsx, TienDo_mau.xlsx cạnh script
  python make_and_check.py check FILE.xlsx   -> so file kế hoạch AI xuất ra với đáp án chuẩn
  python make_and_check.py selftest          -> dựng kế hoạch chuẩn, ghi ra file rồi tự kiểm tra (không gọi AI)

FILE.xlsx là file tải về từ trang xem trước ("Tải file Excel AI đã lập") hoặc file pending_*.xlsx.
"""
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent

# (STT, nội dung, ĐVT, khối lượng, đơn giá)
BOQ = [
    ("A", "PHẦN MÓNG", None, None, None),
    (1, "Đào đất hố móng bằng máy", "m3", 1250.5, 85_000),
    (2, "Bê tông lót móng M100", "m3", 42, 1_350_000),
    (3, "Bê tông móng M300", "m3", 420, 1_650_000),
    (4, "Cốt thép móng", "kg", 31_500, 18_500),
    ("B", "PHẦN THÂN", None, None, None),
    (5, "Xây tường gạch 200", "m3", 310, 1_850_000),
    (6, "Bê tông sàn M300", "m3", 185, 1_700_000),
    (7, "Lắp dựng kết cấu thép mái", "tấn", 24.6, 42_000_000),
    ("C", "HOÀN THIỆN", None, None, None),
    (8, "Trát tường", "m2", 2480, 95_000),
    (9, "Sơn nước", "m2", 2480, 62_000),
    (10, "Điện chiếu sáng", "hệ thống", 1, 180_000_000),     # bẫy: không có công việc trong tiến độ
]
GROUP_TOTALS = {"A": 1_438_742_500, "B": 1_921_200_000, "C": 569_360_000}
GRAND_TOTAL = 3_929_302_500

# (WBS, tên, số ngày, bắt đầu, kết thúc, dự trữ, là tổng hợp)
D = date
SCHEDULE = [
    ("1", "Phần móng", 20, D(2026, 11, 2), D(2026, 11, 21), None, True),
    ("1.1", "Đào đất hố móng", 5, D(2026, 11, 2), D(2026, 11, 6), 0, False),
    ("1.2", "Bê tông lót móng", 3, D(2026, 11, 7), D(2026, 11, 9), 0, False),
    ("1.3", "Cốt thép, ván khuôn, đổ bê tông móng", 12, D(2026, 11, 10), D(2026, 11, 21), 2, False),
    ("2", "Phần thân", 35, D(2026, 11, 22), D(2026, 12, 26), None, True),
    ("2.1", "Xây tường gạch", 15, D(2026, 11, 22), D(2026, 12, 6), 0, False),
    ("2.2", "Đổ bê tông sàn", 10, D(2026, 12, 7), D(2026, 12, 16), 1, False),
    ("2.3", "Lắp dựng kết cấu thép mái", 10, D(2026, 12, 17), D(2026, 12, 26), 0, False),
    ("3", "Hoàn thiện", 17, D(2026, 12, 27), D(2027, 1, 12), None, True),
    ("3.1", "Trát tường", 8, D(2026, 12, 27), D(2027, 1, 3), 0, False),
    ("3.2", "Sơn nước", 7, D(2027, 1, 4), None, 0, False),            # bẫy: thiếu ngày kết thúc (=2027-01-10)
    ("3.3", "Vệ sinh bàn giao", 2, D(2027, 1, 11), D(2027, 1, 12), 0, False),   # bẫy: không có dòng BOQ
]
EXPECTED_END_32 = "2027-01-10"

# đáp án chuẩn cho từng công việc chi tiết: WBS -> (khối lượng chấp nhận, đơn vị, BOQ STT)
TASK_QTY = {
    "1.1": ({1250.5}, "m3", {"1"}),
    "1.2": ({42}, "m3", {"2"}),
    "1.3": ({420, 31_500}, None, {"3", "4"}),        # gộp 2 dòng BOQ khác đơn vị
    "2.1": ({310}, "m3", {"5"}),
    "2.2": ({185}, "m3", {"6"}),
    "2.3": ({24.6}, "tấn", {"7"}),
    "3.1": ({2480}, "m2", {"8"}),
    "3.2": ({2480}, "m2", {"9"}),
    "3.3": (set(), None, set()),
}


# ---- tạo file đầu vào --------------------------------------------------------------------------

def make(out_dir: Path = HERE) -> tuple[Path, Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append(["BẢNG KHỐI LƯỢNG - NHÀ XƯỞNG MẪU (DỮ LIỆU THỬ)"])
    ws.append(["Chủ đầu tư: Công ty TNHH Mẫu Bình Dương   |   Nhà thầu: AZB   |   Gói thầu: Xây dựng nhà xưởng"])
    ws.append([])
    ws.append(["STT", "Nội dung công việc", "ĐVT", "Khối lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"])
    group = None
    for stt, name, unit, qty, price in BOQ:
        if unit is None:
            if group:
                ws.append(["", f"Cộng {group}", None, None, None, GROUP_TOTALS[group]])
            group = stt
            ws.append([stt, name])
        else:
            ws.append([stt, name, unit, qty, price, round(qty * price)])
    ws.append(["", f"Cộng {group}", None, None, None, GROUP_TOTALS[group]])
    ws.append(["", "TỔNG CỘNG", None, None, None, GRAND_TOTAL])
    boq = out_dir / "BOQ_mau.xlsx"
    wb.save(boq)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tiến độ"
    ws.append(["TIẾN ĐỘ THI CÔNG - NHÀ XƯỞNG MẪU (DỮ LIỆU THỬ)"])
    ws.append([])
    ws.append(["WBS", "Tên công việc", "Số ngày", "Bắt đầu", "Kết thúc", "Dự trữ (ngày)"])
    for wbs, name, days, start, end, free, _s in SCHEDULE:
        ws.append([wbs, name, days, datetime.combine(start, datetime.min.time()),
                   datetime.combine(end, datetime.min.time()) if end else None, free])
    for row in ws.iter_rows(min_row=4, min_col=4, max_col=5):
        for c in row:
            c.number_format = "dd/mm/yyyy"
    sched = out_dir / "TienDo_mau.xlsx"
    wb.save(sched)
    return boq, sched


# ---- kiểm tra file AI xuất ra --------------------------------------------------------------------

def _rows(wb, sheet):
    if sheet not in wb.sheetnames:
        return []
    return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def check(path: Path) -> int:
    wb = load_workbook(path, data_only=True)
    ok, bad = [], []

    def expect(cond, label, detail=""):
        (ok if cond else bad).append(f"{label}" + (f"  -> {detail}" if detail and not cond else ""))

    # --- TongQuan
    info = {str(r[0]): r[1] for r in wb["TongQuan"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    expect(bool(info.get("Tên dự án")), "Có tên dự án", str(info.get("Tên dự án")))
    expect("Mẫu" in str(info.get("Chủ đầu tư") or ""), "Chủ đầu tư lấy đúng từ BOQ", str(info.get("Chủ đầu tư")))

    # --- TienDo
    rows = _rows(wb, "TienDo")
    expect(len(rows) == len(SCHEDULE), f"Đủ {len(SCHEDULE)} công việc", f"có {len(rows)}")
    ids = [r[0] for r in rows]
    expect(ids == list(range(1, len(rows) + 1)), "ID chạy 1..n theo thứ tự", str(ids))
    by_wbs = {}
    for r in rows:
        m = re.search(r"WBS\s+([\d.]+)", str(r[8] or ""))
        if m:
            by_wbs[m.group(1)] = r
    for wbs, name, days, start, end, free, is_sum in SCHEDULE:
        r = by_wbs.get(wbs)
        if r is None:
            expect(False, f"[{wbs}] {name}: có dòng với 'WBS {wbs}' trong ghi chú")
            continue
        _id, _cat, title, kind, d, s, e, fl, notes = r[:9]
        want_end = end.isoformat() if end else EXPECTED_END_32
        expect(str(title).strip().lower() == name.lower(), f"[{wbs}] tên công việc", f"'{title}' ≠ '{name}'")
        expect(kind == ("Tổng hợp" if is_sum else "Chi tiết"), f"[{wbs}] loại", str(kind))
        expect(_f(d) == days, f"[{wbs}] số ngày = {days}", str(d))
        expect(str(s) == start.isoformat(), f"[{wbs}] bắt đầu = {start}", str(s))
        expect(str(e) == want_end, f"[{wbs}] kết thúc = {want_end}", str(e))
        if free is not None and not is_sum:
            expect(_f(fl) == free, f"[{wbs}] dự trữ = {free}", str(fl))
        if not is_sum:
            qtys, unit, refs = TASK_QTY[wbs]
            kl = re.search(r"KL:\s*([\d.,]+)\s*([^·]*)", str(notes or ""))
            if qtys:
                got = _f(kl.group(1).replace(",", "")) if kl else None
                expect(got in qtys, f"[{wbs}] khối lượng ∈ {sorted(qtys)}", f"{got}")
                if unit and kl:
                    expect(kl.group(2).strip().lower() == unit, f"[{wbs}] đơn vị = {unit}", kl.group(2).strip())
                got_refs = set(re.findall(r"[\w.]+", (re.search(r"BOQ:\s*([^·]*)", str(notes or "")) or [None, ""])[1]))
                expect(refs <= got_refs, f"[{wbs}] ghép BOQ {sorted(refs)}", f"{sorted(got_refs)}")
            else:
                expect(kl is None, f"[{wbs}] không bịa khối lượng (không có dòng BOQ)", str(notes))
    # cha–con: ngày của dòng tổng hợp phủ các con
    for r in rows:
        if r[3] == "Tổng hợp":
            kids = [k for k in rows if k[3] == "Chi tiết" and k[1] == r[2]]
            if kids:
                expect(str(r[5]) == min(str(k[5]) for k in kids) and str(r[6]) == max(str(k[6]) for k in kids),
                       f"Tổng hợp '{r[2]}' phủ ngày các con",
                       f"{r[5]}→{r[6]} vs {min(str(k[5]) for k in kids)}→{max(str(k[6]) for k in kids)}")

    # --- ChiPhi
    costs = _rows(wb, "ChiPhi")
    expect(len(costs) == 10, "10 khoản chi phí (không có dòng Cộng / TỔNG / tiêu đề nhóm)", f"có {len(costs)}")
    expect(not any(re.match(r"\s*(cộng|tổng)", str(c[1] or ""), re.I) for c in costs), "Không lẫn dòng tổng")
    by_name = {str(c[1]).strip().lower(): c for c in costs}
    total = 0.0
    for stt, name, unit, qty, price in BOQ:
        if unit is None:
            continue
        c = by_name.get(name.lower())
        if c is None:
            expect(False, f"[BOQ {stt}] {name}: có trong ChiPhi")
            continue
        amount = round(qty * price)
        total += _f(c[6]) or 0
        expect(str(c[3]).lower() == unit.lower(), f"[BOQ {stt}] ĐVT = {unit}", str(c[3]))
        expect(_f(c[4]) == qty, f"[BOQ {stt}] khối lượng = {qty}", str(c[4]))
        expect(_f(c[5]) == price, f"[BOQ {stt}] đơn giá = {price:,}", str(c[5]))
        expect(_f(c[6]) == amount, f"[BOQ {stt}] thành tiền = {amount:,}", str(c[6]))
    expect(round(total) == GRAND_TOTAL, f"Tổng thành tiền = {GRAND_TOTAL:,}", f"{total:,.0f}")

    # --- cảnh báo (bẫy)
    warns = " | ".join(str(r[0]) for r in _rows(wb, "GhiChuAI")).lower()
    expect("điện" in warns, "Cảnh báo: BOQ 'Điện chiếu sáng' không có trong tiến độ", warns[:120] or "(không có sheet GhiChuAI)")
    expect("vệ sinh" in warns, "Cảnh báo: 'Vệ sinh bàn giao' không có dòng BOQ", warns[:120])
    expect("3.2" in warns or "sơn" in warns, "Cảnh báo: thiếu ngày kết thúc công việc 3.2 (suy ra)", warns[:120])

    for line in ok:
        print("  OK   ", line)
    for line in bad:
        print("  LỖI  ", line)
    print(f"\n{len(ok)} đạt, {len(bad)} lỗi")
    return 1 if bad else 0


# ---- tự kiểm tra bộ kiểm tra (không gọi AI) ------------------------------------------------------

def selftest() -> int:
    sys.path.insert(0, str(HERE.parents[1]))
    from azb_hub.services import ai_planner as ap

    tasks = []
    for i, (wbs, name, days, start, end, free, is_sum) in enumerate(SCHEDULE, 1):
        grp = next(s[1] for s in SCHEDULE if s[6] and wbs.split(".")[0] == s[0])
        qtys, unit, refs = TASK_QTY.get(wbs, (set(), None, set()))
        tasks.append(ap.AiTask(
            id=i, wbs=wbs, title=name, is_summary=is_sum, category=None if is_sum else grp, planned_days=days,
            start_date=start.isoformat(), end_date=(end or date(2027, 1, 10)).isoformat(), float_days=free,
            quantity=min(qtys) if qtys else None, unit=unit, boq_refs=sorted(refs)))
    costs = [ap.AiCost(stt=int(s), name=n, unit=u_, quantity=q, unit_price=p, budget_amount=round(q * p))
             for s, n, u_, q, p in BOQ if u_]
    plan = ap.AiPlan(project=ap.AiProject(name="Nhà xưởng mẫu", client="Công ty TNHH Mẫu Bình Dương",
                                          contractor="AZB"),
                     tasks=tasks, costs=costs, resources=[], methods=[], risks=[],
                     warnings=["BOQ 10 'Điện chiếu sáng' không có trong tiến độ.",
                               "Công việc 3.3 'Vệ sinh bàn giao' không có dòng BOQ.",
                               "Công việc 3.2 thiếu ngày kết thúc, suy ra 2027-01-10 từ thời lượng 7 ngày."])
    out = HERE / "_selftest_ai_output.xlsx"
    out.write_bytes(ap.build_workbook(plan))
    try:
        return check(out)
    finally:
        out.unlink(missing_ok=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "make":
        for p in make():
            print("Đã tạo", p)
    elif cmd == "check" and len(sys.argv) > 2:
        sys.exit(check(Path(sys.argv[2])))
    elif cmd == "selftest":
        sys.exit(selftest())
    else:
        print(__doc__)
